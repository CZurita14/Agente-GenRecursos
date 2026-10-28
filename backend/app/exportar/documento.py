"""Documentos estáticos sin recursos externos ni interpretación de HTML recibido."""
from io import BytesIO
from pathlib import Path
import re
from threading import Lock
from xml.sax.saxutils import escape

from app.exportar.html import NOMBRES, PIE
from app.exportar.presentacion import es_quiz, presentar


class FormulaInvalida(ValueError):
    pass


_lock = Lock()  # Matplotlib y el registro de fuentes no son reentrantes.
_matematica = re.compile(r'(\$\$.*?\$\$|\$[^$\n]+\$|\\\(.*?\\\)|\\\[.*?\\\])', re.S)


def _fragmentos(texto):
    for parte in _matematica.split(str(texto)):
        if not parte:
            continue
        if _matematica.fullmatch(parte):
            expr = parte[2:-2] if parte.startswith(('$$', '\\(', '\\[')) else parte[1:-1]
            if len(expr) > 500 or expr.count('{') > 40 or re.search(r'\\(?:input|include|href|url|write|def)\b', expr):
                raise FormulaInvalida('Fórmula no admitida')
            try:
                from matplotlib.mathtext import math_to_image
                from matplotlib import rc_context
                from PIL import Image
                salida = BytesIO()
                with _lock, rc_context({'text.usetex': False}):
                    math_to_image('$' + expr + '$', salida, dpi=130, format='png')
                salida.seek(0)
                with Image.open(salida) as im:
                    ancho, alto = im.size
                if ancho > 4000 or alto > 1500:
                    raise FormulaInvalida('Fórmula demasiado grande')
                salida.seek(0)
                yield salida, (ancho, alto)
            except FormulaInvalida:
                raise
            except (ValueError, RuntimeError, RecursionError):
                raise FormulaInvalida('No se pudo representar una fórmula') from None
        else:
            yield parte, None


def bloques_documento(recurso, contexto):
    """Proyección explícita: nunca exportar evidencia o razonamiento interno."""
    yield {'tipo': 'texto', 'titulo': NOMBRES.get(recurso['herramienta'], recurso['herramienta']),
           'texto': contexto.asignatura + ' · ' + contexto.momento_titulo}
    contenido = recurso['contenido']
    if es_quiz(contenido):
        for n, item in enumerate(contenido, 1):
            yield {'tipo': 'grupo', 'titulo': '', 'mantener': True, 'bloques': [
                {'tipo': 'texto', 'titulo': f'Pregunta {n}', 'texto': item.get('contexto', '') + '\n' + item['pregunta']},
                {'tipo': 'lista', 'titulo': '', 'items': [f'{chr(65+k)}) {o["texto"]}' for k, o in enumerate(item['opciones'])]}]}
        yield {'tipo': 'salto'}
        yield {'tipo': 'texto', 'titulo': 'Respuestas y justificaciones', 'texto': ''}
        for n, item in enumerate(contenido, 1):
            o = item['opciones'][item['correcta']]
            yield {'tipo': 'texto', 'titulo': f'Pregunta {n}: {chr(65+item["correcta"])}',
                   'texto': o['texto'] + '\n' + o.get('justificacion', '')}
    else:
        yield from presentar(recurso['herramienta'], contenido)
    yield {'tipo': 'texto', 'titulo': 'Datos del recurso', 'texto':
           f'{contexto.asignatura} · {contexto.clase_id} · {contexto.fecha:%d/%m/%Y}\n'
           + 'Resultados de aprendizaje: ' + '; '.join(f'{r.id}: {r.texto}' for r in contexto.ra)
           + '\nPolítica de IA: ' + (contexto.politica_ia or 'No indicada en el sílabo') + '\n' + PIE}


def exportar_documento(recurso, contexto, formato):
    if formato not in ('pdf', 'docx'):
        raise ValueError('Formato no soportado')
    return _docx(list(bloques_documento(recurso, contexto))) if formato == 'docx' else _pdf(list(bloques_documento(recurso, contexto)))


def _docx(bloques):
    from docx import Document
    from docx.shared import Inches, Cm, Pt, RGBColor
    d = Document()
    d.sections[0].page_width, d.sections[0].page_height = Cm(21), Cm(29.7)
    d.sections[0].top_margin = d.sections[0].bottom_margin = Cm(2)
    for nombre in ('Normal', 'List Paragraph', 'List Bullet', 'Heading 1', 'Heading 2', 'Heading 3'):
        estilo = d.styles[nombre]
        estilo.font.name = 'Arial'
        estilo.font.size = Pt(11 if not nombre.startswith('Heading') else 14)
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.space_before = Pt(8 if nombre.startswith('Heading') else 0)
        estilo.paragraph_format.keep_together = True
    def texto(p, valor):
        for parte, tam in _fragmentos(valor):
            if tam:
                run = p.add_run()
                imagen = run.add_picture(parte, width=Inches(min(tam[0]/130, 6)))
                imagen._inline.docPr.set('descr', 'Fórmula matemática: ' + str(valor))
            else:
                p.add_run(parte)
    def bloque(b, nivel=1):
        if b['tipo'] == 'salto':
            d.add_page_break()
            return
        if b.get('titulo'):
            texto(d.add_heading('', min(nivel, 9)), b['titulo'])
        if b['tipo'] == 'texto':
            texto(d.add_paragraph(), b['texto'])
        elif b['tipo'] == 'lista':
            for i, x in enumerate(b['items'], 1):
                texto(d.add_paragraph(style='List Paragraph' if b.get('numerada') else 'List Bullet'), f'{i}. {x}' if b.get('numerada') else x)
        elif b['tipo'] == 'tabla':
            from docx.oxml import OxmlElement
            t = d.add_table(rows=1, cols=len(b['columnas']))
            t.style = 'Table Grid'
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for c, x in zip(t.rows[0].cells, b['columnas']):
                texto(c.paragraphs[0], x)
            for fila in b['filas']:
                for c, x in zip(t.add_row().cells, fila):
                    texto(c.paragraphs[0], x)
        elif b['tipo'] == 'grupo':
            inicio = len(d.paragraphs)
            for hijo in b['bloques']:
                bloque(hijo, nivel+1)
            if b.get('mantener'):
                for p in d.paragraphs[inicio:-1]:
                    p.paragraph_format.keep_with_next = True
    for b in bloques:
        bloque(b)
    salida = BytesIO()
    d.save(salida)
    return salida.getvalue()


def _pdf(bloques):
    from matplotlib import get_data_path
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Image, Table, TableStyle, PageBreak, KeepTogether
    with _lock:
        if 'Recurso' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('Recurso', str(Path(get_data_path()) / 'fonts/ttf/DejaVuSans.ttf')))
    estilos = getSampleStyleSheet()
    for estilo in estilos.byName.values():
        estilo.fontName = 'Recurso'
    historia = []
    def texto(valor, estilo='BodyText', ancho=450):
        partes = []
        for parte, tam in _fragmentos(valor):
            if tam:
                escala = min(1, ancho/(tam[0]*72/130))
                partes.append(Image(parte, width=tam[0]*72/130*escala, height=tam[1]*72/130*escala))
            else:
                partes.append(Paragraph(escape(parte).replace('\n', '<br/>'), estilos[estilo]))
        return partes or [Paragraph('', estilos[estilo])]
    def bloque(b, nivel=1):
        if b['tipo'] == 'salto':
            historia.append(PageBreak())
            return
        if b.get('titulo'):
            historia.extend(texto(b['titulo'], f'Heading{min(nivel, 3)}'))
        if b['tipo'] == 'texto':
            historia.extend(texto(b['texto']))
        elif b['tipo'] == 'lista':
            for i, x in enumerate(b['items'], 1):
                historia.extend(texto((f'{i}. ' if b.get('numerada') else '• ') + x))
        elif b['tipo'] == 'tabla':
            filas = [[texto(x, ancho=450/len(b['columnas'])-12) for x in f] for f in [b['columnas'], *b['filas']]]
            t = Table(filas, colWidths=[450/len(b['columnas'])]*len(b['columnas']), repeatRows=1, splitInRow=1)
            t.setStyle(TableStyle([('GRID', (0,0), (-1,-1), .5, '#888888'), ('VALIGN', (0,0), (-1,-1), 'TOP'), ('BACKGROUND', (0,0), (-1,0), '#eeeeee')]))
            historia.append(t)
        elif b['tipo'] == 'grupo':
            inicio = len(historia)
            for hijo in b['bloques']:
                bloque(hijo, nivel+1)
            if b.get('mantener'):
                grupo = historia[inicio:]
                historia[inicio:] = [KeepTogether(grupo)]
    for b in bloques:
        bloque(b)
    salida = BytesIO()
    SimpleDocTemplate(salida, pagesize=A4, leftMargin=60, rightMargin=60).build(historia)
    return salida.getvalue()
